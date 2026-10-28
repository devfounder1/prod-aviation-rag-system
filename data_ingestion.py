"""
Модуль для загрузки, очистки и разбиения (чанкинга) технических документов.
"""
import logging
from pathlib import Path
from typing import List
from docx import Document
import PyPDF2 
import pdfplumber
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document as LangChainDocument
from config import settings

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s : %(message)s')

class DocumentIngestor:
    """
    Класс-оркестратор для обработки сырых документов.
    """
    def __init__(self, documents_path: str = None):
        """
        Инициализация обработчика документов.
        
        Args:
            documents_path: Путь к папке с документами (опционально). 
                           Если не указан, берётся из settings.knowledge_base_path
        """
        # Если путь передан явно — используем его, иначе берём из настроек
        if documents_path:
            self.documents_path = Path(documents_path)
        else:
            self.documents_path = Path(settings.knowledge_base_path) if settings.knowledge_base_path else Path(__file__).parent / "documents_ae"
        
        if not self.documents_path.exists():
            raise FileNotFoundError(
                f"Папка с документами не найдена: {self.documents_path}\n"
                f"Укажите путь в настройках приложения."
            )
        
        # Инициализируем сплиттер (разбиватель текста)
        self.splitter = RecursiveCharacterTextSplitter(
            chunk_size = settings.chunk_size,
            chunk_overlap = settings.chunk_overlap,
            separators=["\n\n", "\n", ". ", ", ", " ", ""],
            length_function = len
        )
    
    def load_txt(self, file_path: Path) -> str:
        """Чтение обычного текстового файла"""
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
        
    def load_docx(self, file_path: Path) -> str:
        """
        Чтение Word-документа с извлечением текста и таблиц в Markdown.
        """
        doc = Document(file_path)
        text_parts = []
        
        # Извлекаем текст из параграфов
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
                
        # Извлекаем таблицы и конвертируем в Markdown
        for table_idx, table in enumerate(doc.tables):
            text_parts.append(f"\n--- ТАБЛИЦА {table_idx + 1} (Markdown формат) ---")
            markdown_lines = []
            
            for i, row in enumerate(table.rows):
                # Очищаем текст ячейки от лишних переносов строк, чтобы не ломать Markdown
                cells = [str(cell.text).strip().replace('\n', ' ') for cell in row.cells]
                row_str = "| " + " | ".join(cells) + " |"
                markdown_lines.append(row_str)
                
                if i == 0:
                    # Добавляем разделитель заголовка таблицы (требование стандарта Markdown)
                    separator = "|" + "|".join(["---" for _ in cells]) + "|"
                    markdown_lines.append(separator)
            
            text_parts.append("\n".join(markdown_lines))
            
        return '\n'.join(text_parts)
    
    def load_pdf(self, file_path: Path) -> str:
        """Чтение PDF-файла с улучленным извлечением таблиц."""
        text = ""
        try:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    # 1. Извлекаем обычный текст
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + '\n'
                    
                    # Извлекаем таблицы и конвертируем в Markdown
                    tables = page.extract_tables()
                    for table in tables:
                        if table:
                            text += "\n--- ТАБЛИЦА (Markdown формат) ---\n"
                            # Превращаем таблицу в Markdown-формат, который идеально понимает LLM
                            markdown_table = "| " + " | ".join([str(cell or "").strip() for cell in table[0]]) + " |\n"
                            markdown_table += "| " + " | ".join(["---" for _ in table[0]]) + " |\n"
                            for row in table[1:]:
                                markdown_table += "| " + " | ".join([str(cell or "").strip() for cell in row]) + " |\n"
                            text += markdown_table + "\n"
        except Exception as e:
            logging.warning(f"pdfplumber не справился с {file_path.name}, используем запасной вариант: {e}")
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + '\n'
        
        return text
    
    def process(self) -> List[LangChainDocument]:
        """
        Главный метод: сканирует папку, читает файлы, разбивает на чанки 
        и возвращает список объектов LangChain Document с метаданными.
        """
        logging.info(f"Начало сканирования папки : {self.documents_path.absolute()}")
        
        supported_text = {'.txt', '.docx', '.pdf'}
        langchain_docs = []
        
        # Список паттернов для файлов, которые НЕ нужно индексировать
        skip_patterns = ['содержание', 'contents', 'оглавление', 'table_of_contents']
        
        for file_path in self.documents_path.rglob('*'):
            # Проверяем, что это файл и его расширение нам подходит
            if file_path.is_file() and file_path.suffix.lower() in supported_text:
                
                if file_path.name.startswith('~$'):
                    continue
                
                # ПРОПУСКАЕМ ФАЙЛЫ-ОГЛАВЛЕНИЯ
                file_name_lower = file_path.name.lower()
                if any(pattern in file_name_lower for pattern in skip_patterns):
                    logging.info(f"Пропущен файл-оглавление: {file_path.name}")
                    continue
                
                try: 
                    if file_path.suffix.lower() == '.txt':
                        text = self.load_txt(file_path)
                    elif file_path.suffix.lower() == '.docx':
                        text = self.load_docx(file_path)
                    else : 
                        text = self.load_pdf(file_path)
                    
                    # Разбиваем текст на чанки
                    chunks = self.splitter.split_text(text)
                    
                    # Для каждого чанка создаем объект LangChain Document с метаданными
                    for i, chunk in enumerate(chunks):
                        if not chunk.strip():
                            continue
                        
                        langchain_docs.append(LangChainDocument(
                            page_content=chunk,
                            metadata ={
                                "source": str(file_path.absolute()),
                                "filename": file_path.name,
                                "chunk_index" : i,
                                "chunk_total": len(chunks)
                            }
                        ))
                
                except Exception as e: 
                    logging.error(f"Ошибка обработки файла {file_path.name}: {e}")
        
        logging.info(f"Успешно обработано и разбито на {len(langchain_docs)} чанков.")
        return langchain_docs

if __name__ == "__main__" : 
    print("=" * 60)
    print("ТЕСТ ЗАГРУЗКИ ДОКУМЕНТОВ")
    print("=" * 60)
    
    ingestor = DocumentIngestor()
    docs = ingestor.process()
    
    if docs:
        print(f"\nПример первого чанка:")
        print(f"Файл: {docs[0].metadata['filename']}")
        print(f"Часть: {docs[0].metadata['chunk_index'] + 1} из {docs[0].metadata['chunk_total']}")
        print(f"Длина: {len(docs[0].page_content)} символов")
        print(f"Текст:\n{'-'*40}\n{docs[0].page_content[:400]}...\n{'-'*40}")
    else:
        print("Документы не найдены или не обработаны.")