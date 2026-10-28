"""Загрузка пользовательских документов: парсинг -> чанки -> e5-small -> ChromaDB."""
import re
import shutil
import time
from pathlib import Path

import chromadb
import pandas as pd
from config import settings  # <-- ИСПОЛЬЗУЕМ НАШУ СОВРЕМЕННУЮ КОНФИГУРАЦИЮ

BASE_DIR = Path(__file__).resolve().parent
SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".csv", ".xlsx"}

# ==========================================
# ПАРСИНГ ФАЙЛОВ
# ==========================================
def _read_text_bytes(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "cp1251", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")

def _df_to_text(df, sheet=None) -> str:
    lines = []
    if sheet:
        lines.append(f"Лист: {sheet}")
    lines.append(" | ".join(str(c) for c in df.columns))
    for _, row in df.iterrows():
        lines.append(" | ".join("" if pd.isna(v) else str(v) for v in row))
    return "\n".join(lines)

def parse_file(path: Path) -> str:
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"Формат {ext} не поддерживается. Можно: {', '.join(sorted(SUPPORTED))}")

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        pages = []
        for i, page in enumerate(reader.pages, 1):
            pages.append(f"[страница {i}]\n{page.extract_text() or ''}")
        return "\n\n".join(pages)

    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n\n".join(parts)

    if ext in (".txt", ".md"):
        return _read_text_bytes(path)

    if ext == ".csv":
        for enc in ("utf-8-sig", "cp1251", "latin-1"):
            for sep in (",", ";", "\t"):
                try:
                    df = pd.read_csv(path, encoding=enc, sep=sep)
                    if df.shape[1] > 1:
                        return _df_to_text(df)
                except Exception:
                    continue
        df = pd.read_csv(path, encoding="utf-8", sep=None, engine="python")
        return _df_to_text(df)

    if ext == ".xlsx":
        sheets = pd.read_excel(path, sheet_name=None)
        return "\n\n".join(_df_to_text(df, name) for name, df in sheets.items())

    raise ValueError(f"Не знаю, как читать {ext}")

# ==========================================
# ЧАНКИНГ
# ==========================================
def chunk_text(text: str, size=800, overlap=150):
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks, cur = [], ""
    for p in paragraphs:
        if len(cur) + len(p) + 2 <= size:
            cur = (cur + "\n\n" + p).strip()
        else:
            if cur:
                chunks.append(cur)
            while len(p) > size:
                cut = p.rfind(". ", 0, size)
                if cut < size // 2:
                    cut = size
                chunks.append(p[:cut + 1].strip())
                p = p[cut + 1:].strip()
            cur = p
    if cur:
        chunks.append(cur)
    out = []
    for i, c in enumerate(chunks):
        out.append((chunks[i - 1][-overlap:] + " … " + c) if i > 0 and overlap else c)
    return out

# ==========================================
# ЭМБЕДДИНГИ
# ==========================================
class _Embedder:
    _model = None

    @classmethod
    def model(cls):
        if cls._model is None:
            import torch
            from sentence_transformers import SentenceTransformer
            name = settings.embedder_model_name
            device = settings.embedder_device
            cls._model = SentenceTransformer(name, device=device, local_files_only=True)
            print(f"[DOCS] Эмбеддинг-модель {name} на {device}", flush=True)
        return cls._model

    @classmethod
    def embed_docs(cls, texts):
        vecs = cls.model().encode([f"passage: {t}" for t in texts],
                                  normalize_embeddings=True, batch_size=32,
                                  show_progress_bar=False)
        return vecs.tolist()

    @classmethod
    def embed_query(cls, q):
        vec = cls.model().encode([f"query: {q}"], normalize_embeddings=True)
        return vec.tolist()[0]

# ==========================================
# ХРАНИЛИЩЕ ChromaDB (ИЗОЛИРОВАННОЕ)
# ==========================================
class UserDocStore:
    def __init__(self):
        # ИСПОЛЬЗУЕМ ОТДЕЛЬНУЮ ПАПКУ, ЧТОБЫ НЕ ПУТАТЬ С PRL-27
        db_path = Path(settings.chroma_persist_dir).parent / "user_docs_chroma"
        db_path.mkdir(parents=True, exist_ok=True)
        self.client = chromadb.PersistentClient(path=str(db_path))
        self.coll = self.client.get_or_create_collection(
            name=settings.user_docs_collection,
            metadata={"hnsw:space": "cosine"},
        )

    def count(self):
        return self.coll.count()

    def doc_ids(self):
        res = self.coll.get(include=["metadatas"])
        return sorted({m["doc_id"] for m in res["metadatas"]}) if res["metadatas"] else []

    def enforce_limit(self, chat_id, max_docs):
        res = self.coll.get(where={"chat_id": chat_id}, include=["metadatas"])
        if not res["ids"]:
            return
        
        ids_in_chat = sorted(list({m["doc_id"] for m in res["metadatas"]}))
        while len(ids_in_chat) >= max_docs:
            old = ids_in_chat.pop(0)
            self.coll.delete(where={"doc_id": old})
            print(f"[DOCS] Лимит чата: старый документ {old} удалён", flush=True)

    def add(self, doc_id, filename, chunks, embeddings, chat_id, category):
        self.coll.add(
            ids=[f"{doc_id}_{i}" for i in range(len(chunks))],
            embeddings=embeddings,
            documents=chunks,
            metadatas=[{
                "doc_id": doc_id, 
                "filename": filename, 
                "chunk": i,
                "uploaded_at": int(time.time()),
                "chat_id": chat_id, 
                "pending": 1,
                "category": category
            } for i in range(len(chunks))],
        )

    def query(self, vector, k, where_clause=None):
        if where_clause:
            return self.coll.query(query_embeddings=[vector], n_results=k, where=where_clause,
                                   include=["documents", "metadatas", "distances"])
        return self.coll.query(query_embeddings=[vector], n_results=k,
                               include=["documents", "metadatas", "distances"])

    def list_docs(self, chat_id=None, pending_only=False):
        if chat_id and pending_only:
            where = {"$and": [{"chat_id": chat_id}, {"pending": 1}]}
        elif chat_id:
            where = {"chat_id": chat_id}
        elif pending_only:
            where = {"pending": 1}
        else:
            where = None
            
        res = self.coll.get(where=where, include=["metadatas"]) if where else self.coll.get(include=["metadatas"])
        docs = {}
        for m in res["metadatas"]:
            d = docs.setdefault(m["doc_id"], {
                "doc_id": m["doc_id"], "filename": m["filename"], "chunks": 0,
                "uploaded_at": m.get("uploaded_at", 0),
                "chat_id": m.get("chat_id"), "pending": m.get("pending", 0)})
            d["chunks"] += 1
        return sorted(docs.values(), key=lambda x: x["uploaded_at"])

# ==========================================
# ПУБЛИЧНОЕ API
# ==========================================
def add_document(path, chat_id="global", category="general", progress=print):
    path = Path(path)
    progress(f"[DOCS] Парсинг {path.name}…")
    text = parse_file(path)
    if not text.strip():
        raise ValueError("Документ пустой или текст не распознан")

    docs_dir = Path(settings.documents_dir)
    docs_dir.mkdir(parents=True, exist_ok=True)
    dest = docs_dir / path.name
    if dest.exists():
        dest = docs_dir / f"{path.stem}_{int(time.time())}{path.suffix}"
    shutil.copy2(path, dest)

    progress(f"[DOCS] Чанкинг ({len(text)} символов)…")
    chunks = chunk_text(text, settings.user_docs_chunk_size, settings.user_docs_chunk_overlap)

    store = UserDocStore()
    store.enforce_limit(chat_id, settings.user_docs_max_docs)

    progress(f"[DOCS] Эмбеддинги для {len(chunks)} фрагментов…")
    emb = _Embedder.embed_docs(chunks)

    doc_id = f"doc_{int(time.time() * 1000)}"
    store.add(doc_id, dest.name, chunks, emb, chat_id, category)
    
    progress(f"[DOCS] Готово: {dest.name} -> {len(chunks)} фрагментов (категория: {category})")
    return {"doc_id": doc_id, "filename": dest.name, "chunks": len(chunks), "category": category}

def retrieve_user_context(query, k=None, chat_id=None):
    k = k or settings.user_docs_top_k
    store = UserDocStore()
    if store.count() == 0:
        return ""
    
    vec = _Embedder.embed_query(query)
    where_clause = {"chat_id": chat_id} if chat_id else None
    
    filename_match = re.search(r'([\w\-\sА-Яа-я]+\.(?:pdf|docx|txt|md|csv|xlsx))', query, re.IGNORECASE)
    
    if filename_match:
        target_file = filename_match.group(1).strip()
        where_file = {"filename": target_file}
        if chat_id:
            where_file = {"$and": [{"chat_id": chat_id}, {"filename": target_file}]}
        
        res = store.query(vec, k, where_file)
        if res["documents"][0] and len(res["documents"][0]) > 0:
            parts = []
            for doc, meta, dist in zip(res["documents"][0], res["metadatas"][0], res["distances"][0]):
                parts.append(f"[Источник: {meta['filename']}]\n{doc}")
            return "\n\n".join(parts)
    
    res = store.query(vec, k, where_clause)
    
    parts = []
    for doc, meta, dist in zip(res["documents"][0], res["metadatas"][0], res["distances"][0]):
        if dist is not None and dist > 0.75:
            continue
        parts.append(f"[Источник: {meta['filename']}]\n{doc}")
    return "\n\n".join(parts)

def list_documents(chat_id=None, pending_only=False):
    store = UserDocStore()
    return store.list_docs(chat_id, pending_only)

def consume_documents(chat_id):
    store = UserDocStore()
    res = store.coll.get(where={"$and": [{"chat_id": chat_id}, {"pending": 1}]}, include=["metadatas"])
    if not res["ids"]:
        return []
    new_metas = []
    for m in res["metadatas"]:
        m2 = dict(m)
        m2["pending"] = 0
        new_metas.append(m2)
    store.coll.update(ids=res["ids"], metadatas=new_metas)
    return sorted({m["filename"] for m in res["metadatas"]})

def remove_document(doc_id: str, delete_file: bool = True) -> bool:
    store = UserDocStore()
    res = store.coll.get(where={"doc_id": doc_id}, include=["metadatas"])
    if not res["ids"]:
        return False
    filename = res["metadatas"][0].get("filename", "")
    store.coll.delete(where={"doc_id": doc_id})
    if delete_file and filename:
        try:
            p = Path(settings.documents_dir) / filename
            if p.exists():
                p.unlink()
        except Exception as e:
            print(f"[DOCS] файл не удалён с диска: {e}", flush=True)
    print(f"[DOCS] Документ удалён: {filename}", flush=True)
    return True